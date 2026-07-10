from openpyxl.utils import get_column_letter

def get_real_column_width(worksheet):
    default_width = worksheet.sheet_format.defaultColWidth or 8.38
    real_width = {}
    for i in range(1, worksheet.max_column + 10):
        real_width[get_column_letter(i)] = default_width
    for col_dim in worksheet.column_dimensions.values():
        if col_dim.width is not None:
            start_idx = col_dim.min if col_dim.min else 1
            end_idx = col_dim.max if col_dim.max else start_idx

            if col_dim.width == 13.0 and col_dim.min is None:
                continue

            for idx in range(start_idx, end_idx+1):
                col_letter = get_column_letter(idx)
                real_width[col_letter] = col_dim.width

    return real_width